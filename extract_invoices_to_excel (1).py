"""
extract_invoices_to_excel.py
=============================

Uses GPT-4o (OpenAI Vision) to read "Vendor Payment Report" invoices
(PDFs, or page screenshots/images) and extract, directly off the page:

    - facility         (e.g. "CMC", "CCWF", "COR", "FSP", "CFP")
    - invoice_number   (e.g. "1260077613")
    - total_amount     (the grand total, e.g. "$120,474.39" -> 120474.39)

It then adds the remaining columns automatically (no LLM call needed):

    - s_no          : 1, 2, 3, ... (sequential)
    - client_name   : fixed value "Astrya Global"
    - file_name     : taken from the SOURCE FILE'S OWN NAME, not from
                       anything printed on the page. E.g. a file named
                       "CMC_Astrya_Invoice_1260077613.pdf" becomes the
                       file_name value "CMC_Astrya Invoice 1260077613"
                       (extension stripped, underscores after the first
                       one turned back into spaces).

Finally, everything is written to a formatted Excel workbook that matches
the target layout (S.No | Client Name | File Name | Invoice number |
Facility | Total Amount) with a bold header row and a Grand Total row.

--------------------------------------------------------------------------
INPUT TYPES SUPPORTED
--------------------------------------------------------------------------
1) PDFs (recommended) - each PDF is one invoice/report, and may have
   multiple pages (facility/invoice number on page 1, grand total on the
   last page, as in the CMC sample). Pages are rasterized to images with
   PyMuPDF before being sent to GPT-4o Vision.

2) Pre-rendered page images (PNG/JPG) - either one image per invoice, or
   multiple images per invoice described in a JSON "groups" file.

--------------------------------------------------------------------------
SETUP
--------------------------------------------------------------------------
    pip install openai pandas openpyxl pymupdf

    export OPENAI_API_KEY="sk-..."

--------------------------------------------------------------------------
USAGE
--------------------------------------------------------------------------
1) A folder of invoice PDFs (each file = one invoice, any number of pages):

    python extract_invoices_to_excel.py --pdf_dir ./invoices \
        --output Astrya_Invoices.xlsx

2) A folder of single-page invoice images:

    python extract_invoices_to_excel.py --images_dir ./invoice_pages \
        --output Astrya_Invoices.xlsx

3) Multi-page image invoices, groups described explicitly:

    groups.json:
    [
      ["invoice_pages/CMC_Astrya_Invoice_1260077613_p1.png",
       "invoice_pages/CMC_Astrya_Invoice_1260077613_p2.png"],
      ["invoice_pages/CCWF_Astrya_Invoice_1260076844_p1.png"]
    ]

    python extract_invoices_to_excel.py --groups_json groups.json \
        --output Astrya_Invoices.xlsx

   (file_name for each group is derived from the FIRST image's filename
   in that group, so name the files after the source invoice.)
"""

import argparse
import base64
import json
import os
import sys
import tempfile
import time
from typing import List, Dict, Optional

import pandas as pd
from openai import OpenAI
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Config
# --------------------------------------------------------------------------

MODEL = "gpt-4o"  # vision-capable model
CLIENT_NAME_DEFAULT = "Astrya Global"
MAX_RETRIES = 3

client = OpenAI()  # reads OPENAI_API_KEY from environment


# --------------------------------------------------------------------------
# File-name handling (NOT from GPT — derived from the source file itself)
# --------------------------------------------------------------------------

def derive_file_name(source_path: str) -> Optional[str]:
    """Turn a source file's own name into the File Name column value.

    'CMC_Astrya_Invoice_1260077613.pdf' -> 'CMC_Astrya Invoice 1260077613'

    Keeps the first underscore (between facility code and the rest),
    and turns any remaining underscores into spaces. Strips the
    extension. If the name doesn't follow the underscore pattern, it's
    returned as-is (extension stripped).
    """
    if not source_path:
        return None
    base = os.path.splitext(os.path.basename(source_path))[0]
    parts = base.split("_")
    if len(parts) >= 2:
        return parts[0] + "_" + " ".join(parts[1:])
    return base


# --------------------------------------------------------------------------
# PDF -> page images
# --------------------------------------------------------------------------

def pdf_to_page_images(pdf_path: str, dpi: int = 200) -> List[str]:
    """Rasterize every page of a PDF to a PNG so it can be sent to the
    GPT-4o Vision API. Requires `pip install pymupdf`."""
    import fitz  # PyMuPDF

    doc = fitz.open(pdf_path)
    tmp_dir = tempfile.mkdtemp(prefix="invoice_pages_")
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)

    out_paths = []
    stem = os.path.splitext(os.path.basename(pdf_path))[0]
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=matrix)
        out_path = os.path.join(tmp_dir, f"{stem}_p{i + 1}.png")
        pix.save(out_path)
        out_paths.append(out_path)
    doc.close()
    return out_paths


# --------------------------------------------------------------------------
# GPT-4o Vision extraction
# --------------------------------------------------------------------------

def encode_image_to_data_url(image_path: str) -> str:
    """Read a local image file and return a base64 data URL for the API."""
    ext = os.path.splitext(image_path)[1].lower().lstrip(".")
    mime = "png" if ext == "png" else "jpeg"
    with open(image_path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode("utf-8")
    return f"data:image/{mime};base64,{b64}"


EXTRACTION_PROMPT = """You are looking at one or more pages of the SAME
"Vendor Payment Report" for a facility. These pages may show:
  - A "Facility" code or name near the top-left of the page. This title usually follows the format "<Code> - <Full Name>". 
    Your task is to extract strictly the first continuous word or code that appears BEFORE the hyphen. 
    For example, if it says "XYZ - Some Facility Name", extract only "XYZ".
  - A line that reads something like:
    "View: Agency = Astrya Global And InvoiceNumber is equal to 1260077613"
    — extract the number that comes after "InvoiceNumber is equal to".
  - A "Report Totals" row near the bottom-right, in a column labeled
    "Amount", showing the grand total in dollars (e.g. "$120,474.39").
    This is the FINAL total for the whole report — not a per-registrant
    subtotal like "Totals for (Registrant): ...".

Extract ONLY the following three fields, exactly as printed on the
page(s), and return STRICT JSON with exactly these keys and no others:

{
  "facility": "<the short facility code, e.g. CMC>",
  "invoice_number": "<the invoice number as digits only, e.g. 1260077613>",
  "total_amount": <the report's grand Total amount as a plain number,
                    e.g. 120474.39 - no $ sign, no commas>
}

Rules:
- Extract these values exactly as they appear on the page — do not
  invent or reformat them.
- If a field cannot be found in the provided pages, set its value to null.
- total_amount must be a JSON number (not a string), with no currency
  symbol or thousands separators.
- invoice_number must be a JSON string of digits only.
- Return ONLY the JSON object. No markdown fences, no commentary.
"""


def call_gpt_vision(image_paths: List[str]) -> Dict:
    """Send one or more pages of a single report to GPT-4o and get back
    the parsed JSON fields. Retries a few times on transient/parse errors."""

    content = [{"type": "text", "text": EXTRACTION_PROMPT}]
    for p in image_paths:
        content.append(
            {
                "type": "image_url",
                "image_url": {"url": encode_image_to_data_url(p)},
            }
        )

    last_err = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = client.chat.completions.create(
                model=MODEL,
                temperature=0,
                response_format={"type": "json_object"},
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "You extract structured data from scanned "
                            "vendor payment report images and reply with "
                            "strict JSON only."
                        ),
                    },
                    {"role": "user", "content": content},
                ],
            )
            raw = response.choices[0].message.content
            
            try:
                import sys, os
                _cp = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
                if _cp not in sys.path: sys.path.insert(0, _cp)
                from core.universal_token_monitor import track_usage as _tm
                
                _tm(
                    response.usage, 
                    model=MODEL, 
                    poc_name="INVOICE_EXCEL",
                    file_name=image_paths[0] if image_paths else "unknown", 
                    step_name="extraction"
                )
            except Exception:
                pass
                
            data = json.loads(raw)
            return {
                "facility": data.get("facility"),
                "invoice_number": data.get("invoice_number"),
                "total_amount": data.get("total_amount"),
            }
        except Exception as e:  # noqa: BLE001 - want to retry on any failure
            last_err = e
            print(f"  [warn] attempt {attempt} failed for "
                  f"{image_paths}: {e}", file=sys.stderr)
            time.sleep(1.5 * attempt)

    print(f"  [error] giving up on {image_paths}: {last_err}",
          file=sys.stderr)
    return {"facility": None, "invoice_number": None, "total_amount": None}


# --------------------------------------------------------------------------
# Extraction pipeline
# --------------------------------------------------------------------------

def extract_all(groups: List[Dict]) -> List[Dict]:
    """groups: list of dicts, each: {"source_name": <path used for the
    File Name column>, "images": [page image paths...]}."""

    records = []
    for i, group in enumerate(groups, start=1):
        images = group["images"]
        source_name = group["source_name"]
        print(f"[{i}/{len(groups)}] Extracting from: {images} "
              f"(file name source: {source_name})")

        fields = call_gpt_vision(images)

        record = {
            "s_no": i,
            "Client_name": CLIENT_NAME_DEFAULT,
            "File_name": derive_file_name(source_name),
            "Invoice_number": fields.get("invoice_number"),
            "Facility": fields.get("facility"),
            "Total_amount": fields.get("total_amount"),
        }
        records.append(record)

    return records


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------

def write_excel(records: List[Dict], output_path: str) -> None:
    df = pd.DataFrame(records, columns=[
        "s_no", "Client_name", "File_name",
        "Invoice_number", "Facility", "Total_amount",
    ])
    df.columns = ["S.No", "Client Name", "File Name",
                  "Invoice number", "Facility", "Total Amount"]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Invoices", startrow=0)

        ws = writer.sheets["Invoices"]

        # --- header styling ---
        header_fill = PatternFill("solid", fgColor="FCE4D6")
        header_font = Font(bold=True)
        thin = Side(style="thin", color="B7B7B7")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        # --- body borders + currency format on Total Amount ---
        amount_col = df.columns.get_loc("Total Amount") + 1
        last_row = len(df) + 1  # +1 for header
        for row_idx in range(2, last_row + 1):
            for col_idx in range(1, len(df.columns) + 1):
                ws.cell(row=row_idx, column=col_idx).border = border
            amt_cell = ws.cell(row=row_idx, column=amount_col)
            amt_cell.number_format = '$#,##0.00'

        # --- grand total row ---
        total_row = last_row + 1
        ws.cell(row=total_row, column=amount_col - 1).value = "Total"
        ws.cell(row=total_row, column=amount_col - 1).font = Font(bold=True)
        total_cell = ws.cell(row=total_row, column=amount_col)
        col_letter = get_column_letter(amount_col)
        total_cell.value = f"=SUM({col_letter}2:{col_letter}{last_row})"
        total_cell.number_format = '$#,##0.00'
        total_cell.font = Font(bold=True)
        total_cell.fill = PatternFill("solid", fgColor="D9E1F2")

        # --- column widths ---
        widths = {"A": 6, "B": 16, "C": 32, "D": 18, "E": 12, "F": 16}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

    print(f"Saved: {output_path}")


# --------------------------------------------------------------------------
# CLI / group discovery
# --------------------------------------------------------------------------

def discover_groups_from_pdfs(pdf_dir: str) -> List[Dict]:
    files = sorted(
        os.path.join(pdf_dir, f)
        for f in os.listdir(pdf_dir)
        if f.lower().endswith(".pdf")
    )
    groups = []
    for pdf_path in files:
        images = pdf_to_page_images(pdf_path)
        groups.append({"source_name": pdf_path, "images": images})
    return groups


def discover_groups_from_images_dir(images_dir: str) -> List[Dict]:
    files = sorted(
        os.path.join(images_dir, f)
        for f in os.listdir(images_dir)
        if f.lower().endswith((".png", ".jpg", ".jpeg"))
    )
    return [{"source_name": f, "images": [f]} for f in files]


def normalize_groups_json(raw_groups: List[List[str]]) -> List[Dict]:
    """groups.json holds plain lists of image paths per invoice; use the
    first image's filename as the File Name source for that invoice."""
    return [
        {"source_name": group[0], "images": group}
        for group in raw_groups
    ]


def main():
    parser = argparse.ArgumentParser(
        description="Extract invoice fields via GPT-4o Vision and write a "
                    "formatted Excel file."
    )
    parser.add_argument("--pdf_dir", help="Folder of invoice PDFs (one per invoice)")
    parser.add_argument("--images_dir", help="Folder of single-page invoice images")
    parser.add_argument("--groups_json", help="JSON file listing multi-page invoice image groups")
    parser.add_argument("--output", default="Astrya_Invoices.xlsx",
                        help="Path to the output .xlsx file")
    args = parser.parse_args()

    provided = [a for a in (args.pdf_dir, args.images_dir, args.groups_json) if a]
    if len(provided) != 1:
        parser.error("Provide exactly one of --pdf_dir, --images_dir, --groups_json")

    if args.pdf_dir:
        groups = discover_groups_from_pdfs(args.pdf_dir)
    elif args.groups_json:
        with open(args.groups_json, "r") as f:
            raw_groups = json.load(f)
        groups = normalize_groups_json(raw_groups)
    else:
        groups = discover_groups_from_images_dir(args.images_dir)

    if not groups:
        print("No invoices found.", file=sys.stderr)
        sys.exit(1)

    records = extract_all(groups)
    write_excel(records, args.output)


if __name__ == "__main__":
    main()
